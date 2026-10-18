# tests/test_export_word.py

"""Экспорт в Word: заливка ячеек по флагу дежурства на дому.

Цвета должны совпадать с HTML и Excel:
  - будний день без дежурства — без заливки;
  - выходной или будний день с дежурством — средний серый (#CCCCCC);
  - выходной с дежурством — тёмно-серый (#969696).
Значения days при этом не меняются и не получают суффиксов.
"""

import pytest

from docx import Document
from docx.oxml.ns import qn

from models.project_model import ProjectModel
from controllers.export_controller import write_word_file


# Январь 2025: 1 — среда (будень), 4 — суббота, 5 — воскресенье.
WEEKDAY_DAY = 1        # столбец c_idx = 2 + (1 - 1) = 2
DUTY_WEEKDAY_DAY = 2   # день 2 — четверг (будень), поставим дежурство
WEEKEND_DUTY_DAY = 4   # суббота, с дежурством
WEEKEND_PLAIN_DAY = 5  # воскресенье, без дежурства


def _cell_fill(cell):
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return None
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        return None
    return shd.get(qn('w:fill'))


@pytest.fixture
def doc_model():
    m = ProjectModel(month=0, year=2025)
    m.add_employee("Иванов И.И.", "permanent")
    emp = m.employees[0]
    emp["days"][WEEKDAY_DAY - 1] = "З"
    emp["days"][DUTY_WEEKDAY_DAY - 1] = "З"
    emp["days"][WEEKEND_DUTY_DAY - 1] = "Д"
    emp["days"][WEEKEND_PLAIN_DAY - 1] = "вых"
    emp["duty_home"][DUTY_WEEKDAY_DAY - 1] = True
    emp["duty_home"][WEEKEND_DUTY_DAY - 1] = True
    return m


def test_weekday_without_duty_has_no_fill(tmp_path, doc_model):
    path = tmp_path / "schedule.docx"
    write_word_file(doc_model, str(path))
    table = Document(str(path)).tables[0]
    emp_cells = table.rows[1].cells
    assert _cell_fill(emp_cells[2 + (WEEKDAY_DAY - 1)]) is None


def test_weekday_with_duty_is_shaded(tmp_path, doc_model):
    path = tmp_path / "schedule.docx"
    write_word_file(doc_model, str(path))
    table = Document(str(path)).tables[0]
    emp_cells = table.rows[1].cells
    assert _cell_fill(emp_cells[2 + (DUTY_WEEKDAY_DAY - 1)]) == "CCCCCC"


def test_weekend_with_duty_is_dark(tmp_path, doc_model):
    path = tmp_path / "schedule.docx"
    write_word_file(doc_model, str(path))
    table = Document(str(path)).tables[0]
    emp_cells = table.rows[1].cells
    assert _cell_fill(emp_cells[2 + (WEEKEND_DUTY_DAY - 1)]) == "969696"


def test_weekend_without_duty_is_grey(tmp_path, doc_model):
    path = tmp_path / "schedule.docx"
    write_word_file(doc_model, str(path))
    table = Document(str(path)).tables[0]
    emp_cells = table.rows[1].cells
    assert _cell_fill(emp_cells[2 + (WEEKEND_PLAIN_DAY - 1)]) == "CCCCCC"


def test_days_values_unchanged_by_duty_flag(tmp_path, doc_model):
    path = tmp_path / "schedule.docx"
    write_word_file(doc_model, str(path))
    table = Document(str(path)).tables[0]
    emp_cells = table.rows[1].cells
    assert emp_cells[2 + (DUTY_WEEKDAY_DAY - 1)].text == "З"
    assert emp_cells[2 + (WEEKEND_DUTY_DAY - 1)].text == "Д"


def test_landscape_orientation(tmp_path, doc_model):
    from docx.enum.section import WD_ORIENT
    path = tmp_path / "schedule.docx"
    write_word_file(doc_model, str(path))
    section = Document(str(path)).sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width > section.page_height


def test_all_employees_and_service_rows_present(tmp_path, doc_model):
    doc_model.add_employee("Петров П.П.", "parttime")
    doc_model.set_note(2, "отпуск")
    path = tmp_path / "schedule.docx"
    write_word_file(doc_model, str(path))
    table = Document(str(path)).tables[0]
    names = [table.rows[1].cells[1].text, table.rows[2].cells[1].text]
    assert "Иванов И.И." in names
    assert "Петров П.П." in names
    # Служебные строки Э/О и примечания — последние три строки
    labels = [table.rows[-3].cells[1].text, table.rows[-2].cells[1].text]
    assert labels == ["Э", "О"]
    assert table.rows[-1].cells[2 + 2].text == "отпуск"
