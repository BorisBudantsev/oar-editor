# controllers/export_controller.py

import calendar

from PyQt5.QtWidgets import QFileDialog, QMessageBox
from PyQt5.QtGui import QTextDocument
from PyQt5.QtPrintSupport import QPrinter

from utils.constants import MONTHS, CATEGORY_PERMANENT, CATEGORY_PARTTIME


# ----------------------------------------------------------------------
# Сбор данных
# ----------------------------------------------------------------------
def _collect_data(model):
    days = model.days_in_month()
    employees = model.employees
    permanent = [e for e in employees if e['category'] == CATEGORY_PERMANENT]
    parttime = [e for e in employees if e['category'] == CATEGORY_PARTTIME]
    ordered = permanent + parttime

    weekend_days = set()
    for day in range(1, days + 1):
        wd = calendar.weekday(model.year, model.month + 1, day)
        if wd in (5, 6):
            weekend_days.add(day)

    return {
        'month': model.month,
        'year': model.year,
        'days': days,
        'employees': ordered,
        'special': model.special,
        'last_permanent_idx': len(permanent) - 1 if permanent and parttime else -1,
        'weekend_days': weekend_days,
    }


def _default_name(model, ext):
    return f"График_{MONTHS[model.month]}_{model.year}.{ext}"


# ----------------------------------------------------------------------
# HTML
# ----------------------------------------------------------------------
def build_html(model, scale=8):
    data = _collect_data(model)
    days = data['days']
    weekend = data['weekend_days']

    css = f"""
    <style>
    body {{ font-family: Arial, sans-serif; font-size: {scale}pt; margin: 10px; }}
    h1 {{ font-size: {scale + 4}pt; margin: 4px 0 10px 0; }}
    table {{ border-collapse: collapse; }}
    th, td {{ border: 1px solid #555; padding: 2px 4px;
              text-align: center; font-size: {scale}pt; }}
    th {{ background-color: #f0f0f0; font-weight: bold; }}
    td.name {{ text-align: left; padding-left: 6px; white-space: nowrap; }}
    td.num {{ width: 30px; }}
    .weekend {{ background-color: #dbdbdb; }}
    .border-after td {{ border-bottom: 3px solid #333; }}
    .special-label {{ font-weight: bold; }}
    </style>
    """

    m = MONTHS[data['month']]
    y = data['year']
    title = f"График работы ОАР на {m} {y}"

    parts = [f"<html><head><meta charset='utf-8'>{css}</head><body>"]
    parts.append(f"<h1>{title}</h1>")
    parts.append("<table>")

    parts.append("<thead><tr>")
    parts.append("<th>№</th><th>ФИО</th>")
    for d in range(1, days + 1):
        cls = "weekend" if d in weekend else ""
        parts.append(f"<th class='{cls}'>{d}</th>")
    parts.append("</tr></thead><tbody>")

    for idx, emp in enumerate(data['employees']):
        row_cls = "border-after" if idx == data['last_permanent_idx'] else ""
        parts.append(f"<tr class='{row_cls}'>")
        parts.append(f"<td class='num'>{idx + 1}</td>")
        parts.append(f"<td class='name'>{emp['name']}</td>")
        for day in range(days):
            value = emp['days'][day] if day < len(emp['days']) else ""
            cls = "weekend" if (day + 1) in weekend else ""
            parts.append(f"<td class='{cls}'>{value}</td>")
        parts.append("</tr>")

    for key, values in data['special'].items():
        parts.append("<tr>")
        parts.append("<td></td>")
        parts.append(f"<td class='special-label'>{key}</td>")
        for day in range(days):
            value = values[day] if day < len(values) else ""
            cls = "weekend" if (day + 1) in weekend else ""
            parts.append(f"<td class='{cls}'>{value}</td>")
        parts.append("</tr>")

    parts.append("</tbody></table></body></html>")
    return "".join(parts)


def export_html(parent, model, scale=8):
    path, _ = QFileDialog.getSaveFileName(
        parent, "Экспорт в HTML", _default_name(model, "html"),
        "HTML (*.html);;Все файлы (*)"
    )
    if not path:
        return
    try:
        html = build_html(model, scale)
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
        QMessageBox.information(parent, "Готово", f"Файл сохранён:\n{path}")
    except Exception as e:
        QMessageBox.critical(parent, "Ошибка", f"Не удалось сохранить HTML:\n{e}")


# ----------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------
def export_excel(parent, model, scale=8):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        QMessageBox.critical(parent, "Нет библиотеки",
            "Не установлен openpyxl.\nВыполните: pip install openpyxl")
        return

    path, _ = QFileDialog.getSaveFileName(
        parent, "Экспорт в Excel", _default_name(model, "xlsx"),
        "Excel (*.xlsx);;Все файлы (*)"
    )
    if not path:
        return

    try:
        data = _collect_data(model)
        days = data['days']
        weekend = data['weekend_days']

        wb = Workbook()
        ws = wb.active
        ws.title = "График"

        thin = Side(border_style="thin", color="555555")
        thick = Side(border_style="medium", color="333333")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        border_thick_bottom = Border(left=thin, right=thin, top=thin, bottom=thick)
        weekend_fill = PatternFill(start_color="DBDBDB", end_color="DBDBDB", fill_type="solid")
        header_fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center")

        title_cell = ws.cell(row=1, column=1,
                             value=f"График работы ОАР на {MONTHS[data['month']]} {data['year']}")
        title_cell.font = Font(bold=True, size=scale + 2)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=days + 2)

        start_row = 3
        ws.cell(row=start_row, column=1, value="№")
        ws.cell(row=start_row, column=2, value="ФИО")
        for d in range(1, days + 1):
            ws.cell(row=start_row, column=2 + d, value=d)
        for c in range(1, days + 3):
            cell = ws.cell(row=start_row, column=c)
            cell.font = Font(bold=True)
            cell.alignment = center
            cell.border = border
            cell.fill = header_fill
        for d in weekend:
            ws.cell(row=start_row, column=2 + d).fill = weekend_fill

        row = start_row + 1
        for idx, emp in enumerate(data['employees']):
            ws.cell(row=row, column=1, value=idx + 1).alignment = center
            ws.cell(row=row, column=2, value=emp['name']).alignment = left
            for day in range(days):
                value = emp['days'][day] if day < len(emp['days']) else ""
                c = ws.cell(row=row, column=3 + day, value=value)
                c.alignment = center

            for c in range(1, days + 3):
                cell = ws.cell(row=row, column=c)
                if idx == data['last_permanent_idx']:
                    cell.border = border_thick_bottom
                else:
                    cell.border = border
                if (c - 2) in weekend:
                    cell.fill = weekend_fill
            row += 1

        for key, values in data['special'].items():
            label_cell = ws.cell(row=row, column=2, value=key)
            label_cell.font = Font(bold=True)
            label_cell.alignment = center
            for day in range(days):
                value = values[day] if day < len(values) else ""
                c = ws.cell(row=row, column=3 + day, value=value)
                c.alignment = center
            for c in range(1, days + 3):
                cell = ws.cell(row=row, column=c)
                cell.border = border
                if (c - 2) in weekend:
                    cell.fill = weekend_fill
            row += 1

        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 30
        for d in range(1, days + 1):
            ws.column_dimensions[get_column_letter(2 + d)].width = 5

        wb.save(path)
        QMessageBox.information(parent, "Готово", f"Файл сохранён:\n{path}")
    except Exception as e:
        QMessageBox.critical(parent, "Ошибка", f"Не удалось сохранить Excel:\n{e}")


# ----------------------------------------------------------------------
# Word
# ----------------------------------------------------------------------
def export_word(parent, model, scale=8):
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_ORIENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        QMessageBox.critical(parent, "Нет библиотеки",
            "Не установлен python-docx.\nВыполните: pip install python-docx")
        return

    path, _ = QFileDialog.getSaveFileName(
        parent, "Экспорт в Word", _default_name(model, "docx"),
        "Word (*.docx);;Все файлы (*)"
    )
    if not path:
        return

    try:
        data = _collect_data(model)
        days = data['days']
        weekend = data['weekend_days']

        doc = Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(f"График работы ОАР на {MONTHS[data['month']]} {data['year']}")
        run.bold = True
        run.font.size = Pt(scale + 4)

        total_rows = len(data['employees']) + 2 + 1
        total_cols = days + 2
        table = doc.add_table(rows=total_rows, cols=total_cols)
        table.style = 'Table Grid'

        hdr = table.rows[0].cells
        hdr[0].text = "№"
        hdr[1].text = "ФИО"
        for d in range(1, days + 1):
            hdr[1 + d].text = str(d)

        for idx, emp in enumerate(data['employees']):
            row = table.rows[idx + 1].cells
            row[0].text = str(idx + 1)
            row[1].text = emp['name']
            for day in range(days):
                value = emp['days'][day] if day < len(emp['days']) else ""
                row[2 + day].text = str(value)

        special_start = len(data['employees']) + 1
        for i, (key, values) in enumerate(data['special'].items()):
            row = table.rows[special_start + i].cells
            row[1].text = key
            for day in range(days):
                value = values[day] if day < len(values) else ""
                row[2 + day].text = str(value)

        def set_cell_bg(cell, color_hex):
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), color_hex)
            tcPr.append(shd)

        for r_idx in range(total_rows):
            for c_idx in range(total_cols):
                cell = table.rows[r_idx].cells[c_idx]
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(scale)
                if c_idx >= 2:
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if c_idx >= 2 and (c_idx - 2) in weekend:
                    set_cell_bg(cell, 'DBDBDB')

        doc.save(path)
        QMessageBox.information(parent, "Готово", f"Файл сохранён:\n{path}")
    except Exception as e:
        QMessageBox.critical(parent, "Ошибка", f"Не удалось сохранить Word:\n{e}")


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
def export_pdf(parent, model, scale=8):
    path, _ = QFileDialog.getSaveFileName(
        parent, "Экспорт в PDF", _default_name(model, "pdf"),
        "PDF (*.pdf);;Все файлы (*)"
    )
    if not path:
        return
    try:
        html = build_html(model, scale)
        doc = QTextDocument()
        doc.setHtml(html)

        printer = QPrinter(QPrinter.HighResolution)
        printer.setOutputFormat(QPrinter.PdfFormat)
        printer.setOutputFileName(path)
        printer.setOrientation(QPrinter.Landscape)
        printer.setPageMargins(10, 10, 10, 10, QPrinter.Millimeter)

        doc.print_(printer)
        QMessageBox.information(parent, "Готово", f"Файл сохранён:\n{path}")
    except Exception as e:
        QMessageBox.critical(parent, "Ошибка", f"Не удалось сохранить PDF:\n{e}")
        # ======================================================================
# ЭКСПОРТ ОТЧЁТА АНАЛИЗА ЗАГРУЖЕННОСТИ
# ======================================================================

def _fmt_cell(v):
    """Пустые значения → '—', иначе строку."""
    if v is None or v == "":
        return "—"
    return str(v)


def _category_ru(cat):
    return "постоянный" if cat == CATEGORY_PERMANENT else "совместитель"


def _default_stats_name(model, ext):
    return f"Анализ_{MONTHS[model.month]}_{model.year}.{ext}"


# ----------------------------------------------------------------------
# HTML
# ----------------------------------------------------------------------
def build_stats_html(model, report, scale=8):
    css = f"""
    <style>
    body {{ font-family: Arial, sans-serif; font-size: {scale}pt; margin: 12px; }}
    h1 {{ font-size: {scale + 4}pt; margin: 4px 0 14px 0; }}
    h2 {{ font-size: {scale + 2}pt; margin: 16px 0 6px 0; border-bottom: 1px solid #999; }}
    table {{ border-collapse: collapse; margin-bottom: 12px; }}
    th, td {{ border: 1px solid #555; padding: 2px 5px;
              text-align: center; font-size: {scale}pt; }}
    th {{ background-color: #f0f0f0; font-weight: bold; }}
    td.name {{ text-align: left; padding-left: 6px; white-space: nowrap; }}
    </style>
    """

    m = MONTHS[model.month]
    y = model.year

    parts = [f"<html><head><meta charset='utf-8'>{css}</head><body>"]
    parts.append(f"<h1>Анализ загруженности ОАР на {m} {y}</h1>")

    # --- Раздел 1: По сотрудникам ---
    parts.append("<h2>По сотрудникам</h2>")
    parts.append("<table>")
    parts.append(
        "<tr><th>ФИО</th><th>Категория</th><th>Рабочих</th><th>Выходных</th>"
        "<th>Д</th><th>Энд</th><th>Опер</th><th>О/Э</th><th>Коэф. занятости</th></tr>"
    )
    for r in report["employees"]:
        parts.append("<tr>")
        parts.append(f"<td class='name'>{r['name']}</td>")
        parts.append(f"<td>{_category_ru(r['category'])}</td>")
        parts.append(f"<td>{r['work']}</td>")
        parts.append(f"<td>{_fmt_cell(r['vyh'])}</td>")
        parts.append(f"<td>{r['d_count']}</td>")
        parts.append(f"<td>{r['endo']}</td>")
        parts.append(f"<td>{r['oper']}</td>")
        parts.append(f"<td>{r['o_e_ratio']}</td>")
        parts.append(f"<td>{_fmt_cell(r['busy_ratio'])}</td>")
        parts.append("</tr>")
    parts.append("</table>")

    # --- Раздел 2: По рабочим местам ---
    parts.append("<h2>По рабочим местам</h2>")
    wp = report["workplaces"]
    blocks = [
        ("Эндоскопия (ЭНД1+ЭНД2+ЭНД3)", wp["endo"]),
        ("Операционная (З+Б+Ж+К)", wp["oper"]),
        ("Дежурства (Д)", wp["d"]),
        ("Отношение Опер/Энд", wp["ratio"]),
        ("Коэффициент занятости (совместители)", wp["busy"]),
    ]
    for title, rows in blocks:
        parts.append(f"<h3 style='font-size:{scale + 1}pt;margin:10px 0 4px 0;'>{title}</h3>")
        parts.append("<table>")
        parts.append("<tr><th>Сотрудник</th><th>Значение</th></tr>")
        for name, val in rows:
            parts.append(f"<tr><td class='name'>{name}</td><td>{val}</td></tr>")
        parts.append("</table>")

    # --- Раздел 3: Общий итог ---
    parts.append("<h2>Общий итог</h2>")
    t = report["totals"]
    parts.append("<table>")
    parts.append(f"<tr><td class='name'>Смен ЭНД1+ЭНД2+ЭНД3</td><td>{t['endo']}</td></tr>")
    parts.append(f"<tr><td class='name'>Смен З+Б+Ж+К</td><td>{t['oper']}</td></tr>")
    parts.append(f"<tr><td class='name'>Смен Ж</td><td>{t['zh']}</td></tr>")
    parts.append("</table>")

    parts.append("</body></html>")
    return "".join(parts)


def export_stats_html(parent, model, report, scale=8):
    path, _ = QFileDialog.getSaveFileName(
        parent, "Экспорт отчёта в HTML", _default_stats_name(model, "html"),
        "HTML (*.html);;Все файлы (*)"
    )
    if not path:
        return
    try:
        html = build_stats_html(model, report, scale)
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
        QMessageBox.information(parent, "Готово", f"Файл сохранён:\n{path}")
    except Exception as e:
        QMessageBox.critical(parent, "Ошибка", f"Не удалось сохранить HTML:\n{e}")


# ----------------------------------------------------------------------
# Excel — три листа
# ----------------------------------------------------------------------
def export_stats_excel(parent, model, report, scale=8):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        QMessageBox.critical(parent, "Нет библиотеки",
            "Не установлен openpyxl.\nВыполните: pip install openpyxl")
        return

    path, _ = QFileDialog.getSaveFileName(
        parent, "Экспорт отчёта в Excel", _default_stats_name(model, "xlsx"),
        "Excel (*.xlsx);;Все файлы (*)"
    )
    if not path:
        return

    try:
        wb = Workbook()

        thin = Side(border_style="thin", color="555555")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center")

        m = MONTHS[model.month]
        y = model.year

        # ---- Лист 1: Сотрудники ----
        ws = wb.active
        ws.title = "Сотрудники"
        headers = ["ФИО", "Категория", "Рабочих", "Выходных",
                   "Д", "Энд", "Опер", "О/Э", "Коэф. занятости"]
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = Font(bold=True)
            cell.alignment = center
            cell.border = border
            cell.fill = header_fill

        for i, r in enumerate(report["employees"], start=2):
            vals = [
                r["name"], _category_ru(r["category"]),
                r["work"], _fmt_cell(r["vyh"]), r["d_count"],
                r["endo"], r["oper"], r["o_e_ratio"],
                _fmt_cell(r["busy_ratio"]),
            ]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=i, column=c, value=v)
                cell.border = border
                cell.alignment = left if c == 1 else center

        ws.column_dimensions['A'].width = 28
        for col in range(2, 10):
            ws.column_dimensions[get_column_letter(col)].width = 14

        # ---- Лист 2: Рабочие места ----
        ws2 = wb.create_sheet("Рабочие места")
        row = 1
        wp = report["workplaces"]
        blocks = [
            ("Эндоскопия (ЭНД1+ЭНД2+ЭНД3)", wp["endo"]),
            ("Операционная (З+Б+Ж+К)", wp["oper"]),
            ("Дежурства (Д)", wp["d"]),
            ("Отношение Опер/Энд", wp["ratio"]),
            ("Коэффициент занятости (совместители)", wp["busy"]),
        ]
        for title, rows in blocks:
            cell = ws2.cell(row=row, column=1, value=title)
            cell.font = Font(bold=True, size=scale + 2)
            ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            row += 1

            for c, h in enumerate(["Сотрудник", "Значение"], start=1):
                cell = ws2.cell(row=row, column=c, value=h)
                cell.font = Font(bold=True)
                cell.alignment = center
                cell.border = border
                cell.fill = header_fill
            row += 1

            for name, val in rows:
                ws2.cell(row=row, column=1, value=name).border = border
                c2 = ws2.cell(row=row, column=2, value=val)
                c2.border = border
                c2.alignment = center
                row += 1

            row += 1  # пустая строка между блоками

        ws2.column_dimensions['A'].width = 30
        ws2.column_dimensions['B'].width = 18

        # ---- Лист 3: Итог ----
        ws3 = wb.create_sheet("Итог")
        t = report["totals"]
        ws3.cell(row=1, column=1, value="Показатель").font = Font(bold=True)
        ws3.cell(row=1, column=2, value="Значение").font = Font(bold=True)
        ws3.cell(row=2, column=1, value="Смен ЭНД1+ЭНД2+ЭНД3")
        ws3.cell(row=2, column=2, value=t['endo']).alignment = center
        ws3.cell(row=3, column=1, value="Смен З+Б+Ж+К")
        ws3.cell(row=3, column=2, value=t['oper']).alignment = center
        ws3.cell(row=4, column=1, value="Смен Ж")
        ws3.cell(row=4, column=2, value=t['zh']).alignment = center
        for r_ in range(1, 5):
            for c_ in range(1, 3):
                ws3.cell(row=r_, column=c_).border = border
        ws3.column_dimensions['A'].width = 30
        ws3.column_dimensions['B'].width = 14

        wb.save(path)
        QMessageBox.information(parent, "Готово", f"Файл сохранён:\n{path}")
    except Exception as e:
        QMessageBox.critical(parent, "Ошибка", f"Не удалось сохранить Excel:\n{e}")


# ----------------------------------------------------------------------
# Word
# ----------------------------------------------------------------------
def export_stats_word(parent, model, report, scale=8):
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_ORIENT
    except ImportError:
        QMessageBox.critical(parent, "Нет библиотеки",
            "Не установлен python-docx.\nВыполните: pip install python-docx")
        return

    path, _ = QFileDialog.getSaveFileName(
        parent, "Экспорт отчёта в Word", _default_stats_name(model, "docx"),
        "Word (*.docx);;Все файлы (*)"
    )
    if not path:
        return

    try:
        doc = Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

        m = MONTHS[model.month]
        y = model.year

        h1 = doc.add_paragraph()
        h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = h1.add_run(f"Анализ загруженности ОАР на {m} {y}")
        run.bold = True
        run.font.size = Pt(scale + 4)

        # ---- Раздел 1: Сотрудники ----
        h2 = doc.add_paragraph()
        r = h2.add_run("По сотрудникам")
        r.bold = True
        r.font.size = Pt(scale + 2)

        headers = ["ФИО", "Категория", "Рабочих", "Выходных",
                   "Д", "Энд", "Опер", "О/Э", "Коэф. занятости"]
        t1 = doc.add_table(rows=1, cols=len(headers))
        t1.style = 'Table Grid'
        for c, h in enumerate(headers):
            cell = t1.rows[0].cells[c]
            cell.text = h
            for p in cell.paragraphs:
                for rr in p.runs:
                    rr.bold = True
                    rr.font.size = Pt(scale)

        for emp in report["employees"]:
            row = t1.add_row().cells
            vals = [
                emp["name"], _category_ru(emp["category"]),
                str(emp["work"]), _fmt_cell(emp["vyh"]), str(emp["d_count"]),
                str(emp["endo"]), str(emp["oper"]), emp["o_e_ratio"],
                _fmt_cell(emp["busy_ratio"]),
            ]
            for i, v in enumerate(vals):
                row[i].text = v
                for p in row[i].paragraphs:
                    for rr in p.runs:
                        rr.font.size = Pt(scale)
                    if i >= 2:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # ---- Раздел 2: Рабочие места ----
        doc.add_paragraph()
        h3 = doc.add_paragraph()
        r = h3.add_run("По рабочим местам")
        r.bold = True
        r.font.size = Pt(scale + 2)

        wp = report["workplaces"]
        blocks = [
            ("Эндоскопия (ЭНД1+ЭНД2+ЭНД3)", wp["endo"]),
            ("Операционная (З+Б+Ж+К)", wp["oper"]),
            ("Дежурства (Д)", wp["d"]),
            ("Отношение Опер/Энд", wp["ratio"]),
            ("Коэффициент занятости (совместители)", wp["busy"]),
        ]
        for title, rows in blocks:
            sub = doc.add_paragraph()
            rr = sub.add_run(title)
            rr.bold = True
            rr.font.size = Pt(scale + 1)

            t = doc.add_table(rows=1, cols=2)
            t.style = 'Table Grid'
            hdr = t.rows[0].cells
            hdr[0].text = "Сотрудник"
            hdr[1].text = "Значение"
            for c_ in hdr:
                for p in c_.paragraphs:
                    for r_ in p.runs:
                        r_.bold = True
                        r_.font.size = Pt(scale)
            for name, val in rows:
                row = t.add_row().cells
                row[0].text = str(name)
                row[1].text = str(val)
                for i, c_ in enumerate(row):
                    for p in c_.paragraphs:
                        for r_ in p.runs:
                            r_.font.size = Pt(scale)
                        if i == 1:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # ---- Раздел 3: Итог ----
        doc.add_paragraph()
        h4 = doc.add_paragraph()
        r = h4.add_run("Общий итог")
        r.bold = True
        r.font.size = Pt(scale + 2)

        t = doc.add_table(rows=1, cols=2)
        t.style = 'Table Grid'
        hdr = t.rows[0].cells
        hdr[0].text = "Показатель"
        hdr[1].text = "Значение"
        for c_ in hdr:
            for p in c_.paragraphs:
                for r_ in p.runs:
                    r_.bold = True
                    r_.font.size = Pt(scale)

        totals = report["totals"]
        for label, val in [
            ("Смен ЭНД1+ЭНД2+ЭНД3", totals['endo']),
            ("Смен З+Б+Ж+К", totals['oper']),
            ("Смен Ж", totals['zh']),
        ]:
            row = t.add_row().cells
            row[0].text = label
            row[1].text = str(val)
            for i, c_ in enumerate(row):
                for p in c_.paragraphs:
                    for r_ in p.runs:
                        r_.font.size = Pt(scale)
                    if i == 1:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.save(path)
        QMessageBox.information(parent, "Готово", f"Файл сохранён:\n{path}")
    except Exception as e:
        QMessageBox.critical(parent, "Ошибка", f"Не удалось сохранить Word:\n{e}")


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
def export_stats_pdf(parent, model, report, scale=8):
    path, _ = QFileDialog.getSaveFileName(
        parent, "Экспорт отчёта в PDF", _default_stats_name(model, "pdf"),
        "PDF (*.pdf);;Все файлы (*)"
    )
    if not path:
        return
    try:
        html = build_stats_html(model, report, scale)
        doc = QTextDocument()
        doc.setHtml(html)

        printer = QPrinter(QPrinter.HighResolution)
        printer.setOutputFormat(QPrinter.PdfFormat)
        printer.setOutputFileName(path)
        printer.setOrientation(QPrinter.Landscape)
        printer.setPageMargins(10, 10, 10, 10, QPrinter.Millimeter)

        doc.print_(printer)
        QMessageBox.information(parent, "Готово", f"Файл сохранён:\n{path}")
    except Exception as e:
        QMessageBox.critical(parent, "Ошибка", f"Не удалось сохранить PDF:\n{e}")